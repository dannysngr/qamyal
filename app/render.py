"""Оформление итога: документ Word и страница с плеером."""
import html
import json
from pathlib import Path

from common import fmt_duration, fmt_time, projects_of


def meta_line(res: dict) -> str:
    ps = projects_of(res)
    parts = [(f"проекты: {', '.join(ps)}" if len(ps) > 1 else f"проект: {ps[0]}") if ps else "",
             res.get("date", ""), fmt_duration(res.get("duration", 0))]
    if res.get("participants"):
        parts.append("участники: " + ", ".join(res["participants"]))
    return " · ".join(p for p in parts if p)


def task_text(t, with_project: bool = False) -> str:
    if isinstance(t, str):
        return t
    s = t.get("what", "")
    if t.get("who"):
        s = f"{t['who']}: {s}"
    if t.get("when"):
        s += f" (срок: {t['when']})"
    if with_project and t.get("project"):
        s += f" — {t['project']}"
    return s


def sections(res: dict) -> list:
    many = len(projects_of(res)) > 1          # у записи несколько проектов — у задачи видно, к какому она
    out = [(title, [task_text(x, many) for x in res.get(key) or []])
           for title, key in (("Кратко", "summary"), ("Решения", "decisions"), ("Задачи", "tasks"))]
    marks = []
    for m in res.get("marks") or []:
        near = min(res["turns"], key=lambda t: abs(t["t"] - m["t"]), default=None)
        what = m.get("note") or (f"{near['who']}: {near['text'][:120]}…" if near else "")
        marks.append(f"{fmt_time(m['t'])} — {what}")
    out.append(("Отмеченные моменты", marks))
    total = sum(x["sec"] for x in res.get("stats") or []) or 1
    out.append(("Кто сколько говорил", [f"{x['who']}: {fmt_duration(x['sec'])} ({x['sec'] * 100 // total}%)"
                                         for x in res.get("stats") or []]))
    return out


def write_docx(res: dict, path: Path) -> None:
    from docx import Document
    from docx.shared import Pt, RGBColor

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(11)
    doc.add_heading(res["title"], level=1)
    p = doc.add_paragraph(meta_line(res))
    p.runs[0].font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    for title, items in sections(res):
        if items:
            doc.add_heading(title, level=2)
            for it in items:
                doc.add_paragraph(it, style="List Bullet")
    doc.add_heading("Расшифровка", level=2)
    for t in res["turns"]:
        p = doc.add_paragraph()
        r = p.add_run(f"[{fmt_time(t['t'])}] ")
        r.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
        r.font.size = Pt(9)
        p.add_run(f"{t['who']}: ").bold = True
        p.add_run(t["text"])
        p.paragraph_format.space_after = Pt(4)
    doc.save(str(path))


PAGE = """<!doctype html>
<html lang="ru"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<title>{title}</title>
<style>
:root{{--bg:#fafaf8;--fg:#1d1d1b;--muted:#6b6b66;--line:#e4e3de;--card:#fff;--accent:#1f5fbf;--now:#fff3c4}}
@media (prefers-color-scheme:dark){{:root{{--bg:#191917;--fg:#ecebe6;--muted:#9a9a93;--line:#33332f;--card:#22221f;--accent:#7fb0ff;--now:#3d3519}}}}
*{{box-sizing:border-box}}
body{{margin:0;background:var(--bg);color:var(--fg);font:16px/1.55 "Segoe UI",system-ui,sans-serif}}
header{{position:sticky;top:0;z-index:1;background:var(--card);border-bottom:1px solid var(--line);padding:12px 16px}}
header h1{{font-size:19px;margin:0 0 2px}} header .meta{{color:var(--muted);font-size:13px;margin-bottom:8px}}
audio{{width:100%;height:36px}}
main{{max-width:860px;margin:0 auto;padding:8px 16px 80px}}
h2{{font-size:16px;margin:24px 0 6px}} ul{{margin:0;padding-left:22px}}
.turn{{display:grid;grid-template-columns:56px 1fr;gap:8px;padding:6px 8px;border-radius:6px}}
.turn.now{{background:var(--now)}}
.ts{{color:var(--accent);font-size:13px;padding-top:2px;cursor:pointer;text-decoration:none;font-variant-numeric:tabular-nums}}
.who{{font-weight:600}}
</style></head><body>
<header><h1>{title}</h1><div class="meta">{meta}</div><audio id="a" controls preload="metadata" src="запись.mp3"></audio></header>
<main>{sections}<h2>Расшифровка</h2>
{turns}
</main>
<script>
const a=document.getElementById('a'),rows=[...document.querySelectorAll('.turn')],ts=rows.map(r=>+r.dataset.t);
document.querySelectorAll('.ts').forEach(x=>x.onclick=e=>{{e.preventDefault();a.currentTime=+x.closest('.turn').dataset.t;a.play();}});
let cur=-1;a.ontimeupdate=()=>{{let i=-1;for(let k=0;k<ts.length&&ts[k]<=a.currentTime+0.3;k++)i=k;
if(i!==cur){{if(cur>=0)rows[cur].classList.remove('now');if(i>=0){{rows[i].classList.add('now');if(!a.paused)rows[i].scrollIntoView({{block:'center',behavior:'smooth'}});}}cur=i;}}}};
</script></body></html>"""


def write_html(res: dict, path: Path) -> None:
    e = html.escape
    secs = "".join(f"<h2>{e(title)}</h2><ul>{''.join(f'<li>{e(i)}</li>' for i in items)}</ul>"
                   for title, items in sections(res) if items)
    turns = "\n".join(
        f'<div class="turn" data-t="{t["t"]:.1f}"><a class="ts" href="#">{fmt_time(t["t"])}</a>'
        f'<div><span class="who">{e(t["who"])}:</span> {e(t["text"])}</div></div>'
        for t in res["turns"])
    path.write_text(PAGE.format(title=e(res["title"]), meta=e(meta_line(res)), sections=secs,
                                turns=turns), encoding="utf-8")


def file_base(res: dict, what: str = "", limit: int = 70) -> str:
    """Имя файла: «2026-10-19 Замечания по тренажёру» (дата разговора + краткое название)."""
    import re
    from common import safe_name
    m = re.match(r"(\d{2})\.(\d{2})\.(\d{4})", res.get("date", ""))
    title = safe_name(res.get("title") or "Расшифровка")
    if len(title) > limit:
        title = title[:limit].rsplit(" ", 1)[0].rstrip(" ,.—-")
    parts = [f"{m[3]}-{m[2]}-{m[1]}" if m else "", f"{what} — {title}" if what else title]
    return " ".join(p for p in parts if p)


def outputs(job_dir: Path) -> dict:
    """Пути к готовым Word и странице записи (старые записи — «Итог.*»)."""
    names = json.loads((job_dir / "черновики" / "файлы.json").read_text(encoding="utf-8")) \
        if (job_dir / "черновики" / "файлы.json").exists() else {}
    return {k: job_dir / names.get(k, f"Итог.{k}") for k in ("docx", "html")}


def render(job_dir: Path, res: dict) -> None:
    participants = []
    for t in res["turns"]:
        if t["who"] not in participants and not t["who"].startswith("Спикер"):
            participants.append(t["who"])
    res["participants"] = participants
    (job_dir / "черновики" / "итог.json").write_text(
        json.dumps(res, ensure_ascii=False, indent=1), encoding="utf-8")
    base = file_base(res)
    new = {"docx": job_dir / f"{base}.docx", "html": job_dir / f"{base}.html"}
    for kind, old in outputs(job_dir).items():      # название поменялось — прежние файлы убираем
        if old != new[kind] and old.exists():
            try:
                old.unlink()
            except OSError:
                pass                                # открыт в Word — оставляем как есть
    write_docx(res, new["docx"])
    write_html(res, new["html"])
    (job_dir / "черновики" / "файлы.json").write_text(
        json.dumps({k: p.name for k, p in new.items()}, ensure_ascii=False), encoding="utf-8")
